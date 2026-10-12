import gradio as gr
import requests
import json
import os
from PyPDF2 import PdfReader
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet


OLLAMA_URL = "http://localhost:11434/api/generate"
MODEL_NAME = "qwen2.5:0.5b"

os.makedirs("exports", exist_ok=True)


# --------------------------------------------------
# EXTRACT TEXT
# --------------------------------------------------

def extract_text(file):

    if file is None:
        return ""

    filename = file.name
    text = ""

    try:

        if filename.lower().endswith(".pdf"):

            reader = PdfReader(filename)

            for page in reader.pages:

                page_text = page.extract_text()

                if page_text:
                    text += page_text + "\n"


        elif filename.lower().endswith(".docx"):

            doc = Document(filename)

            text = "\n".join(
                p.text for p in doc.paragraphs
                if p.text.strip()
            )


        elif filename.lower().endswith(".txt"):

            with open(filename, "r", encoding="utf-8") as f:
                text = f.read()


        else:

            return "Unsupported file type."


        return text.strip()


    except Exception as e:

        return f"Error extracting document: {e}"


# --------------------------------------------------
# OLLAMA
# --------------------------------------------------

def ollama_generate(prompt):

    try:

        response = requests.post(
            OLLAMA_URL,
            json={
                "model": MODEL_NAME,
                "prompt": prompt,
                "temperature": 0.2,
                "stream": True
            },
            stream=True,
            timeout=300
        )

        response.raise_for_status()

        full_output = ""

        for line in response.iter_lines(decode_unicode=True):

            if not line:
                continue

            try:

                data = json.loads(line)

                if "response" in data:
                    full_output += data["response"]

                if data.get("done", False):
                    break

            except json.JSONDecodeError:
                continue

        if full_output.strip():
            return full_output.strip()

        return "No response generated."


    except Exception as e:

        return f"Ollama connection error: {e}"


# --------------------------------------------------
# SUMMARY
# --------------------------------------------------

def generate_summaries(text):

    short_prompt = f"""
Analyze the following document.

Give a very short summary in only 1 or 2 sentences.

Document:
{text}
"""

    medium_prompt = f"""
Analyze the following document.

Give a clear medium-length summary in one paragraph.
Mention the main purpose, important points and conclusion.

Document:
{text}
"""

    detailed_prompt = f"""
Analyze the following document carefully.

Create a detailed section-wise summary.

If it is a legal or contract document, identify sections such as:
- Parties
- Purpose
- Payment Terms
- Responsibilities
- Confidentiality
- Liability
- Termination
- Risks

If it is a literary document, identify:
- Characters
- Plot
- Major Events
- Themes
- Moral Lessons

Use clear headings and bullet points.

Document:
{text}
"""

    short_summary = ollama_generate(short_prompt)

    medium_summary = ollama_generate(medium_prompt)

    detailed_summary = ollama_generate(detailed_prompt)

    return short_summary, medium_summary, detailed_summary


# --------------------------------------------------
# KEYWORDS / CLAUSES
# --------------------------------------------------

def generate_keywords(text):

    prompt = f"""
Analyze the following document.

Extract the most important keywords, topics, clauses and themes.

For a legal document focus on:
Payment Terms, Confidentiality, Liability,
Termination, Responsibilities, Risks,
Disputes and Obligations.

For a literary document focus on:
Characters, Themes, Plot Events,
Important Ideas and Moral Lessons.

Give the result as a clean bullet list.

Document:
{text}
"""

    return ollama_generate(prompt)


# --------------------------------------------------
# PROCESS DOCUMENT
# --------------------------------------------------

def analyze_document(file):

    if file is None:

        return (
            "Please upload a PDF, DOCX or TXT document.",
            "",
            "",
            ""
        )

    text = extract_text(file)

    if not text:

        return (
            "No text could be extracted from the document.",
            "",
            "",
            ""
        )

    if text.startswith("Unsupported"):

        return (
            text,
            "",
            "",
            ""
        )

    # Limit extremely large documents
    text = text[:30000]

    short_summary, medium_summary, detailed_summary = generate_summaries(text)

    keywords = generate_keywords(text)

    return (
        short_summary,
        medium_summary,
        detailed_summary,
        keywords
    )


# --------------------------------------------------
# QUESTION ANSWERING
# --------------------------------------------------

def answer_question(file, question):

    if file is None:

        return "Please upload a document first."

    if not question.strip():

        return "Please enter a question."

    text = extract_text(file)

    if not text:

        return "No text could be extracted."

    text = text[:30000]

    prompt = f"""
You are an AI document analysis assistant.

Answer the user's question using ONLY the information
available in the document.

If the answer is not available in the document,
say:

"The answer is not available in the document."

Document:
{text}

Question:
{question}

Give a clear and concise answer.
"""

    return ollama_generate(prompt)


# --------------------------------------------------
# EXPORT TXT
# --------------------------------------------------

def export_txt(short_summary, medium_summary,
               detailed_summary, keywords):

    filepath = os.path.abspath(
        "exports/document_analysis.txt"
    )

    content = f"""
AI DOCUMENT ANALYZER
====================

SHORT SUMMARY
-------------
{short_summary}


MEDIUM SUMMARY
--------------
{medium_summary}


DETAILED SUMMARY
----------------
{detailed_summary}


KEYWORDS / KEY CLAUSES
----------------------
{keywords}
"""

    with open(filepath, "w", encoding="utf-8") as f:
        f.write(content)

    return filepath


# --------------------------------------------------
# EXPORT DOCX
# --------------------------------------------------

def export_docx(short_summary, medium_summary,
                detailed_summary, keywords):

    filepath = os.path.abspath(
        "exports/document_analysis.docx"
    )

    doc = Document()

    doc.add_heading(
        "AI Document Analysis Report",
        level=0
    )

    doc.add_heading("Short Summary", level=1)
    doc.add_paragraph(short_summary)

    doc.add_heading("Medium Summary", level=1)
    doc.add_paragraph(medium_summary)

    doc.add_heading("Detailed Summary", level=1)
    doc.add_paragraph(detailed_summary)

    doc.add_heading(
        "Keywords / Key Clauses",
        level=1
    )

    doc.add_paragraph(keywords)

    doc.save(filepath)

    return filepath


# --------------------------------------------------
# EXPORT PDF
# --------------------------------------------------

def export_pdf(short_summary, medium_summary,
               detailed_summary, keywords):

    filepath = os.path.abspath(
        "exports/document_analysis.pdf"
    )

    styles = getSampleStyleSheet()

    doc = SimpleDocTemplate(
        filepath,
        pagesize=A4
    )

    elements = []

    elements.append(
        Paragraph(
            "AI Document Analysis Report",
            styles["Title"]
        )
    )

    elements.append(Spacer(1, 20))

    elements.append(
        Paragraph(
            "Short Summary",
            styles["Heading2"]
        )
    )

    elements.append(
        Paragraph(
            short_summary.replace("\n", "<br/>"),
            styles["BodyText"]
        )
    )

    elements.append(Spacer(1, 15))

    elements.append(
        Paragraph(
            "Medium Summary",
            styles["Heading2"]
        )
    )

    elements.append(
        Paragraph(
            medium_summary.replace("\n", "<br/>"),
            styles["BodyText"]
        )
    )

    elements.append(Spacer(1, 15))

    elements.append(
        Paragraph(
            "Detailed Summary",
            styles["Heading2"]
        )
    )

    elements.append(
        Paragraph(
            detailed_summary.replace("\n", "<br/>"),
            styles["BodyText"]
        )
    )

    elements.append(Spacer(1, 15))

    elements.append(
        Paragraph(
            "Keywords / Key Clauses",
            styles["Heading2"]
        )
    )

    elements.append(
        Paragraph(
            keywords.replace("\n", "<br/>"),
            styles["BodyText"]
        )
    )

    doc.build(elements)

    return filepath


# --------------------------------------------------
# EXPORT ALL
# --------------------------------------------------

def export_all(short_summary,
               medium_summary,
               detailed_summary,
               keywords):

    txt_file = export_txt(
        short_summary,
        medium_summary,
        detailed_summary,
        keywords
    )

    docx_file = export_docx(
        short_summary,
        medium_summary,
        detailed_summary,
        keywords
    )

    pdf_file = export_pdf(
        short_summary,
        medium_summary,
        detailed_summary,
        keywords
    )

    return txt_file, docx_file, pdf_file


# --------------------------------------------------
# CLEAR
# --------------------------------------------------

def clear_all():

    return (
        None,
        "",
        "",
        "",
        "",
        "",
        None,
        None,
        None
    )


# --------------------------------------------------
# GUI
# --------------------------------------------------

with gr.Blocks(
    title="AI Document Analyzer"
) as app:

    gr.Markdown(
        """
# 📄 AI Document Analyzer

### Summarize • Query • Analyze • Export

Upload a **PDF, DOCX or TXT** document and use
AI to summarize, analyze and ask questions about it.
"""
    )

    with gr.Row():

        file_input = gr.File(
            label="📁 Upload Document",
            file_types=[
                ".pdf",
                ".docx",
                ".txt"
            ],
            type="filepath"
        )

        analyze_button = gr.Button(
            "🔍 Analyze Document",
            variant="primary"
        )

        clear_button = gr.Button(
            "🗑️ Clear"
        )


    with gr.Tabs():

        # --------------------------------------------------
        # SUMMARY TAB
        # --------------------------------------------------

        with gr.Tab("📝 Summary"):

            gr.Markdown("### Short Summary")

            short_output = gr.Textbox(
                lines=4,
                label="1–2 Sentence Summary"
            )

            gr.Markdown("### Medium Summary")

            medium_output = gr.Textbox(
                lines=8,
                label="Paragraph Summary"
            )

            gr.Markdown("### Detailed Summary")

            detailed_output = gr.Textbox(
                lines=18,
                label="Section-wise Analysis"
            )


        # --------------------------------------------------
        # Q&A TAB
        # --------------------------------------------------

        with gr.Tab("💬 Q&A"):

            question_input = gr.Textbox(
                lines=3,
                placeholder="Ask a question about the uploaded document...",
                label="Your Question"
            )

            question_button = gr.Button(
                "Ask AI",
                variant="primary"
            )

            answer_output = gr.Textbox(
                lines=10,
                label="AI Answer"
            )


        # --------------------------------------------------
        # KEYWORDS TAB
        # --------------------------------------------------

        with gr.Tab("🔑 Keywords / Key Clauses"):

            keywords_output = gr.Textbox(
                lines=20,
                label="Important Keywords, Clauses and Themes"
            )


        # --------------------------------------------------
        # EXPORT TAB
        # --------------------------------------------------

        with gr.Tab("📤 Export & Reporting"):

            gr.Markdown(
                """
### Export your document analysis

The generated report contains:

- Short Summary
- Medium Summary
- Detailed Summary
- Keywords / Key Clauses
"""
            )

            export_button = gr.Button(
                "📥 Generate Reports",
                variant="primary"
            )

            txt_download = gr.File(
                label="TXT Report"
            )

            docx_download = gr.File(
                label="DOCX Report"
            )

            pdf_download = gr.File(
                label="PDF Report"
            )


    # --------------------------------------------------
    # EVENTS
    # --------------------------------------------------

    analyze_button.click(
        fn=analyze_document,
        inputs=file_input,
        outputs=[
            short_output,
            medium_output,
            detailed_output,
            keywords_output
        ]
    )


    question_button.click(
        fn=answer_question,
        inputs=[
            file_input,
            question_input
        ],
        outputs=answer_output
    )


    export_button.click(
        fn=export_all,
        inputs=[
            short_output,
            medium_output,
            detailed_output,
            keywords_output
        ],
        outputs=[
            txt_download,
            docx_download,
            pdf_download
        ]
    )


    clear_button.click(
        fn=clear_all,
        inputs=[],
        outputs=[
            file_input,
            short_output,
            medium_output,
            detailed_output,
            keywords_output,
            answer_output,
            txt_download,
            docx_download,
            pdf_download
        ]
    )


# --------------------------------------------------
# START APPLICATION
# --------------------------------------------------

if __name__ == "__main__":
    app.launch(
        server_name="127.0.0.1",
        server_port=7868,
        inbrowser=True,
        share=True
    )